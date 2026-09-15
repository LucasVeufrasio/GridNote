"""Leitura e gravação dos formatos suportados.

Na grade tudo é texto: cada célula é uma `str`. Ao gravar em Excel, textos que são
números "limpos" (sem zero à esquerda) voltam a ser números.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from xml.sax.saxutils import escape

OPEN_EXTS = (".csv", ".txt", ".xlsx", ".xls")

# chave -> (rótulo, extensão, descrição curta, é editável ao reabrir?)
SAVE_FORMATS = {
    "xlsx": ("Excel", ".xlsx", "Planilha do Excel (recomendado)", True),
    "xls": ("Excel 97-2003", ".xls", "Excel antigo, até 65.536 linhas", True),
    "csv": ("CSV", ".csv", "Texto separado por ; ou ,", True),
    "txt": ("Texto", ".txt", "Texto simples, com ou sem separador", True),
    "docx": ("Word", ".docx", "Documento com a tabela formatada", False),
    "pdf": ("PDF", ".pdf", "Para imprimir ou enviar", False),
}

# Modos de divisão de um .txt
TXT_MODES = {
    "auto": "Detectar automaticamente",
    "\t": "Tabulação",
    ";": "Ponto e vírgula  ( ; )",
    ",": "Vírgula  ( , )",
    "|": "Barra vertical  ( | )",
    "char": "Cada caractere em uma coluna (largura fixa)",
    "line": "Linha inteira em uma única coluna",
}


class FileFormatError(Exception):
    """Erro com mensagem pronta para mostrar ao usuário."""


@dataclass
class Sheet:
    name: str
    rows: list[list[str]]
    header: bool = False


@dataclass
class Loaded:
    sheets: list[Sheet]
    meta: dict = field(default_factory=dict)


# ----------------------------------------------------------------------------
# utilidades
# ----------------------------------------------------------------------------

def rectangular(rows: list[list[str]]) -> list[list[str]]:
    """Garante que todas as linhas tenham o mesmo número de colunas (mínimo 1x1)."""
    if not rows:
        return [[""]]
    width = max((len(r) for r in rows), default=0) or 1
    for r in rows:
        if len(r) < width:
            r.extend([""] * (width - len(r)))
    return rows


def format_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        return format(v, ".15g")
    if isinstance(v, dt.datetime):
        if v.hour == v.minute == v.second == 0:
            return v.strftime("%d/%m/%Y")
        return v.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, dt.time):
        return v.strftime("%H:%M:%S")
    return str(v)


_NUMBER = re.compile(r"-?(0|[1-9]\d{0,14})(\.\d{1,15})?")


def to_excel_value(s: str):
    if s and _NUMBER.fullmatch(s):
        return float(s) if "." in s else int(s)
    return s


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Remove linhas e colunas vazias no fim."""
    while rows and not any(rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if r[i]:
                width = max(width, i + 1)
                break
    return [r[:width] for r in rows]


def _split_lines(text: str) -> list[str]:
    lines = re.split(r"\r\n|\n|\r", text)
    if lines and lines[-1] == "":
        lines.pop()
    return lines


def decode_bytes(raw: bytes) -> tuple[str, str]:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace"), "utf-8-sig"
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace"), "cp1252"


def detect_delimiter(lines: list[str]) -> str | None:
    sample = [l for l in lines[:200] if l.strip()]
    if not sample:
        return None
    for d in ("\t", ";", "|", ","):
        counts = [l.count(d) for l in sample]
        top = max(set(counts), key=counts.count)
        if top > 0 and counts.count(top) >= 0.8 * len(counts):
            return d
    return None


def split_text(lines: list[str], mode: str) -> list[list[str]]:
    if mode == "char":
        return [list(l) for l in lines]
    if mode == "line":
        return [[l] for l in lines]
    return [l.split(mode) for l in lines]


def join_row(row: list[str], mode: str) -> str:
    """Inverso de split_text: devolve a linha de texto que gerou as células."""
    if mode == "char":
        # largura fixa: buracos antes do último caractere viram espaço
        last = max((i for i, v in enumerate(row) if v), default=-1)
        return "".join(v or " " for v in row[: last + 1])
    if mode == "line":
        return "\t".join(row).rstrip("\t") if len(row) > 1 else (row[0] if row else "")
    return ("\t" if mode == "cells" else mode).join(row)


_LOOKS_NUMERIC = re.compile(r"[-+R$\s]*[\d.,/:%]+\s*")


def looks_like_header(rows: list[list[str]]) -> bool:
    """Primeira linha toda preenchida com textos (não números/datas) e com dados abaixo."""
    if len(rows) < 2 or not rows[0]:
        return False
    first = [v.strip() for v in rows[0]]
    return all(first) and not any(_LOOKS_NUMERIC.fullmatch(v) for v in first) and len(set(first)) == len(first)


def resolve_txt_mode(lines: list[str], mode: str) -> str:
    if mode != "auto":
        return mode
    return detect_delimiter(lines) or "line"


# ----------------------------------------------------------------------------
# leitura
# ----------------------------------------------------------------------------

def read_file(path: str | Path, txt_mode: str = "auto") -> Loaded:
    path = Path(path)
    ext = path.suffix.lower()
    try:
        raw = path.read_bytes()
    except OSError as e:
        raise FileFormatError(f"Não foi possível abrir o arquivo:\n{e}") from e

    if ext in (".xlsx", ".xls"):
        return _read_excel_like(path, raw)
    if ext == ".csv":
        return _read_csv(raw)
    if ext == ".txt":
        return _read_txt(raw, txt_mode)
    raise FileFormatError(f"Formato não suportado: {ext}")


def text_lines_of(path: str | Path) -> list[str]:
    """Primeiras linhas de um arquivo de texto (para a pré-visualização)."""
    raw = Path(path).read_bytes()[:256_000]
    text, _ = decode_bytes(raw)
    return _split_lines(text)


def _text_meta(raw: bytes) -> tuple[str, dict]:
    text, enc = decode_bytes(raw)
    newline = "\r\n" if "\r\n" in text[:65536] else "\n"
    return text, {"encoding": enc, "newline": newline}


def _read_txt(raw: bytes, mode: str) -> Loaded:
    text, meta = _text_meta(raw)
    lines = _split_lines(text)
    mode = resolve_txt_mode(lines, mode)
    meta["txt_mode"] = mode
    rows = rectangular(split_text(lines, mode))
    header = mode not in ("char", "line") and looks_like_header(rows)
    return Loaded([Sheet("Planilha1", rows, header)], meta)


def _read_csv(raw: bytes) -> Loaded:
    text, meta = _text_meta(raw)
    lines = _split_lines(text[:256_000])
    delim = detect_delimiter(lines) or ("," if "," in text[:4096] else ";")
    meta["delimiter"] = delim
    rows = rectangular([list(r) for r in csv.reader(io.StringIO(text, newline=""), delimiter=delim)])
    return Loaded([Sheet("Planilha1", rows, looks_like_header(rows))], meta)


def _read_excel_like(path: Path, raw: bytes) -> Loaded:
    # Muitos sistemas exportam ".xls" que na verdade é HTML ou texto; olhamos o conteúdo.
    if raw[:2] == b"PK":
        return _read_xlsx(path)
    if raw[:4] == b"\xd0\xcf\x11\xe0":
        return _read_xls(raw)
    text, meta = _text_meta(raw)
    if "<table" in text[:200_000].lower():
        parser = _HTMLTables()
        parser.feed(text)
        sheets = [Sheet(f"Tabela{i + 1}", rectangular(_trim(t))) for i, t in enumerate(parser.tables) if t]
        if sheets:
            return Loaded(sheets, meta)
    lines = _split_lines(text)
    delim = detect_delimiter(lines)
    if delim:
        return Loaded([Sheet("Planilha1", rectangular(split_text(lines, delim)))], meta)
    raise FileFormatError("Este arquivo não parece ser uma planilha do Excel válida.")


def _read_xlsx(path: Path) -> Loaded:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 - openpyxl levanta vários tipos
        raise FileFormatError(f"Não foi possível ler a planilha:\n{e}") from e
    sheets = []
    try:
        for ws in wb.worksheets:
            rows: list[list[str]] = []
            empty_run = 0
            for values in ws.iter_rows(values_only=True):
                row = [format_value(v) for v in values]
                rows.append(row)
                # algumas planilhas declaram um milhão de linhas vazias
                empty_run = 0 if any(row) else empty_run + 1
                if empty_run > 20_000:
                    break
            rows = rectangular(_trim(rows))
            sheets.append(Sheet(ws.title, rows, looks_like_header(rows)))
    finally:
        wb.close()
    if not sheets:
        raise FileFormatError("A planilha não tem abas com dados.")
    return Loaded(sheets, {})


def _read_xls(raw: bytes) -> Loaded:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=raw)
    except Exception as e:  # noqa: BLE001
        raise FileFormatError(f"Não foi possível ler o arquivo .xls:\n{e}") from e
    sheets = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        row.append(format_value(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                    except Exception:  # noqa: BLE001
                        row.append(format_value(cell.value))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(format_value(bool(cell.value)))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(xlrd.error_text_from_code.get(cell.value, "#ERRO"))
                else:
                    row.append(format_value(cell.value))
            rows.append(row)
        rows = rectangular(_trim(rows))
        sheets.append(Sheet(sh.name, rows, looks_like_header(rows)))
    return Loaded(sheets or [Sheet("Planilha1", [[""]])], {})


class _HTMLTables(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._depth = 0
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._depth += 1
            if self._depth == 1:
                self.tables.append([])
        elif self._depth == 1 and tag == "tr":
            self._row = []
        elif self._depth == 1 and tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag == "table":
            self._depth = max(0, self._depth - 1)
        elif self._depth == 1 and tag in ("td", "th") and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif self._depth == 1 and tag == "tr" and self._row is not None:
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


# ----------------------------------------------------------------------------
# gravação
# ----------------------------------------------------------------------------

def write_file(path: str | Path, fmt: str, sheets: list[Sheet], meta: dict | None = None) -> None:
    meta = meta or {}
    path = Path(path)
    writers = {
        "csv": _write_csv, "txt": _write_txt, "xlsx": _write_xlsx,
        "xls": _write_xls, "docx": _write_docx, "pdf": _write_pdf,
    }
    if fmt not in writers:
        raise FileFormatError(f"Formato de gravação desconhecido: {fmt}")
    # grava num temporário e troca no fim, para nunca deixar o arquivo pela metade
    tmp = path.with_name(f".{path.stem}.gridnote-tmp{path.suffix}")
    try:
        writers[fmt](tmp, sheets, meta)
        tmp.replace(path)
    except PermissionError as e:
        raise FileFormatError(
            "Não foi possível gravar. O arquivo pode estar aberto em outro programa "
            f"(Excel, Word...). Feche-o e tente de novo.\n\n{e}"
        ) from e
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)


def _encoding(meta: dict) -> str:
    return meta.get("encoding") or "utf-8-sig"


def _write_csv(path: Path, sheets: list[Sheet], meta: dict) -> None:
    if meta.get("txt_mode") in ("char", "line"):
        # dividido por caractere/linha na tela: grava o texto original de cada linha
        _write_txt(path, sheets, meta)
        return
    rows = sheets[0].rows
    with open(path, "w", encoding=_encoding(meta), errors="replace", newline="") as f:
        w = csv.writer(f, delimiter=meta.get("delimiter") or ";", lineterminator=meta.get("newline") or "\r\n")
        w.writerows(rows)


def _write_txt(path: Path, sheets: list[Sheet], meta: dict) -> None:
    rows = sheets[0].rows
    mode = meta.get("txt_mode") or "\t"
    out = [join_row(r, mode) for r in rows]
    nl = meta.get("newline") or "\r\n"
    with open(path, "w", encoding=_encoding(meta), errors="replace", newline="") as f:
        f.write(nl.join(out))
        if out:
            f.write(nl)


def _sheet_names(sheets: list[Sheet]) -> list[str]:
    used, names = set(), []
    for s in sheets:
        base = re.sub(r"[\[\]:*?/\\]", "_", s.name).strip("'")[:31] or "Planilha"
        name, i = base, 2
        while name.lower() in used:
            suffix = f" ({i})"
            name = base[: 31 - len(suffix)] + suffix
            i += 1
        used.add(name.lower())
        names.append(name)
    return names


def _write_xlsx(path: Path, sheets: list[Sheet], meta: dict) -> None:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook(write_only=True)
    bold = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F3864")
    for sheet, name in zip(sheets, _sheet_names(sheets)):
        ws = wb.create_sheet(name)
        for i, row in enumerate(sheet.rows):
            out = []
            for v in row:
                v = ILLEGAL_CHARACTERS_RE.sub("", v)
                value = to_excel_value(v)
                if (sheet.header and i == 0) or (isinstance(value, str) and value.startswith("=")):
                    cell = WriteOnlyCell(ws, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"  # nunca vira fórmula
                    if sheet.header and i == 0:
                        cell.font, cell.fill = bold, fill
                    out.append(cell)
                else:
                    out.append(value)
            ws.append(out)
    wb.save(path)


def _write_xls(path: Path, sheets: list[Sheet], meta: dict) -> None:
    import xlwt

    for s in sheets:
        if len(s.rows) > 65536 or (s.rows and len(s.rows[0]) > 256):
            raise FileFormatError(
                f"A aba \"{s.name}\" tem {len(s.rows)} linhas e {len(s.rows[0])} colunas.\n"
                "O formato .xls aceita no máximo 65.536 linhas e 256 colunas. Salve como .xlsx."
            )
    wb = xlwt.Workbook(encoding="utf-8")
    head = xlwt.easyxf("font: bold on, colour white; pattern: pattern solid, fore_colour dark_blue")
    for sheet, name in zip(sheets, _sheet_names(sheets)):
        ws = wb.add_sheet(name, cell_overwrite_ok=True)
        for r, row in enumerate(sheet.rows):
            for c, v in enumerate(row):
                if not v:
                    continue
                if len(v) > 32767:
                    raise FileFormatError(f"Uma célula da linha {r + 1} passa de 32.767 caracteres (limite do .xls).")
                value = to_excel_value(v)
                if sheet.header and r == 0:
                    ws.write(r, c, value, head)
                else:
                    ws.write(r, c, value)
    wb.save(str(path))


_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _docx_runs(text: str, header: bool) -> str:
    props = '<w:rPr><w:b/><w:color w:val="FFFFFF"/><w:sz w:val="18"/></w:rPr>' if header else '<w:rPr><w:sz w:val="18"/></w:rPr>'
    parts = _XML_BAD.sub("", text).split("\n")
    runs = []
    for i, part in enumerate(parts):
        br = "<w:br/>" if i else ""
        runs.append(f'<w:r>{props}{br}<w:t xml:space="preserve">{escape(part)}</w:t></w:r>')
    return "".join(runs)


def _write_docx(path: Path, sheets: list[Sheet], meta: dict) -> None:
    for s in sheets:
        if s.rows and len(s.rows[0]) > 63:
            raise FileFormatError(
                f"A aba \"{s.name}\" tem {len(s.rows[0])} colunas, mas uma tabela do Word aceita no máximo 63.\n"
                "Salve como PDF ou Excel."
            )
    page_w, margin = 16838, 720
    border = '<w:{0} w:val="single" w:sz="4" w:space="0" w:color="8EA9DB"/>'
    borders = "".join(border.format(b) for b in ("top", "left", "bottom", "right", "insideH", "insideV"))
    body = []
    for s in sheets:
        ncols = len(s.rows[0]) if s.rows else 1
        col_w = (page_w - 2 * margin) // ncols
        body.append(
            '<w:p><w:pPr><w:spacing w:after="120"/></w:pPr><w:r><w:rPr><w:b/><w:color w:val="1F3864"/>'
            f'<w:sz w:val="28"/></w:rPr><w:t xml:space="preserve">{escape(_XML_BAD.sub("", s.name))}</w:t></w:r></w:p>'
        )
        body.append(
            '<w:tbl><w:tblPr><w:tblW w:w="5000" w:type="pct"/>'
            f"<w:tblBorders>{borders}</w:tblBorders>"
            '<w:tblCellMar><w:left w:w="60" w:type="dxa"/><w:right w:w="60" w:type="dxa"/></w:tblCellMar>'
            "</w:tblPr><w:tblGrid>" + f'<w:gridCol w:w="{col_w}"/>' * ncols + "</w:tblGrid>"
        )
        for i, row in enumerate(s.rows):
            is_head = s.header and i == 0
            tr = ["<w:tr>", "<w:trPr><w:tblHeader/></w:trPr>" if is_head else ""]
            for v in row:
                shade = '<w:tcPr><w:shd w:val="clear" w:color="auto" w:fill="1F3864"/></w:tcPr>' if is_head else ""
                tr.append(f"<w:tc>{shade}<w:p><w:pPr><w:spacing w:after=\"0\"/></w:pPr>{_docx_runs(v, is_head)}</w:p></w:tc>")
            tr.append("</w:tr>")
            body.append("".join(tr))
        body.append("</w:tbl><w:p/>")
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
        + "".join(body)
        + f'<w:sectPr><w:pgSz w:w="{page_w}" w:h="11906" w:orient="landscape"/>'
        f'<w:pgMar w:top="{margin}" w:right="{margin}" w:bottom="{margin}" w:left="{margin}" '
        'w:header="360" w:footer="360" w:gutter="0"/></w:sectPr></w:body></w:document>'
    )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="word/document.xml"/></Relationships>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", content_types)
        z.writestr("_rels/.rels", rels)
        z.writestr("word/document.xml", document)


def _write_pdf(path: Path, sheets: list[Sheet], meta: dict) -> None:
    # Desenhado com o próprio Qt: não precisa de biblioteca de PDF extra.
    from PySide6.QtCore import QMarginsF, QRectF, Qt
    from PySide6.QtGui import QColor, QFont, QFontMetricsF, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageOrientation(QPageLayout.Orientation.Landscape)
    writer.setPageMargins(QMarginsF(10, 10, 10, 12), QPageLayout.Unit.Millimeter)
    writer.setResolution(150)
    writer.setTitle(meta.get("title") or path.stem)
    writer.setCreator("GridNote")

    painter = QPainter(writer)
    try:
        W, H = writer.width(), writer.height()
        page_no = 0
        for si, sheet in enumerate(sheets):
            rows = sheet.rows
            ncols = len(rows[0]) if rows else 0
            if ncols == 0:
                continue
            size = 8.0
            font = QFont("Segoe UI", 1)
            font.setPointSizeF(size)
            fm = QFontMetricsF(font, writer)
            pad = fm.averageCharWidth() * 1.2

            def natural_widths():
                ws = [fm.horizontalAdvance(str(c + 1)) + 2 * pad for c in range(ncols)]
                for row in rows[:400]:
                    for c, v in enumerate(row):
                        ws[c] = max(ws[c], min(fm.horizontalAdvance(v) + 2 * pad, W * 0.35))
                return ws

            widths = natural_widths()
            if sum(widths) > W:
                size = max(5.0, size * W / sum(widths))
                font.setPointSizeF(size)
                fm = QFontMetricsF(font, writer)
                pad = fm.averageCharWidth() * 1.2
                widths = natural_widths()
            total = sum(widths)
            if total > W:
                widths = [w * W / total for w in widths]
            elif total < W * 0.6:
                widths = [w * (W * 0.6) / total for w in widths]
            rh = fm.height() * 1.6

            title_font = QFont("Segoe UI", 12, QFont.Weight.Bold)
            small = QFont("Segoe UI", 7)
            title_h = QFontMetricsF(title_font, writer).height() * 1.8
            footer_h = QFontMetricsF(small, writer).height() * 1.8
            header_row = rows[0] if sheet.header else None
            body = rows[1:] if sheet.header else rows

            def new_page():
                nonlocal page_no
                if page_no:
                    writer.newPage()
                page_no += 1
                painter.setFont(title_font)
                painter.setPen(QColor("#1F3864"))
                painter.drawText(QRectF(0, 0, W, title_h), Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, sheet.name)
                painter.setFont(small)
                painter.setPen(QColor("#6F7F99"))
                painter.drawText(QRectF(0, H - footer_h, W, footer_h), Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter, f"Página {page_no}")
                y = title_h
                if header_row is not None:
                    draw_row(header_row, y, head=True, alt=False)
                    y += rh
                return y

            def draw_row(row, y, head, alt):
                x = 0.0
                painter.setFont(font)
                for c, v in enumerate(row):
                    rect = QRectF(x, y, widths[c], rh)
                    if head:
                        painter.fillRect(rect, QColor("#1F3864"))
                    elif alt:
                        painter.fillRect(rect, QColor("#EEF3FB"))
                    painter.setPen(QPen(QColor("#B7C6DE"), 1))
                    painter.drawRect(rect)
                    painter.setPen(QColor("#FFFFFF") if head else QColor("#1B2433"))
                    text = fm.elidedText(v.replace("\n", " "), Qt.TextElideMode.ElideRight, max(0.0, widths[c] - 2 * pad))
                    painter.drawText(rect.adjusted(pad, 0, -pad, 0), Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter, text)
                    x += widths[c]

            y = new_page()
            for i, row in enumerate(body):
                if y + rh > H - footer_h:
                    y = new_page()
                draw_row(row, y, head=False, alt=bool(i % 2))
                y += rh
    finally:
        painter.end()
